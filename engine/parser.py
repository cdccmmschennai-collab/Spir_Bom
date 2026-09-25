"""
Parses a SPIR (Spare Parts & Interchangeability Record) Excel file into a
structured, sheet-agnostic form. Handles both single-page SPIRs (one
"MAIN SHEET") and multi-page SPIRs ("MAIN SHEET" + "SHEET (2)", "SHEET (3)"...).

This is deliberately generic: it does NOT hardcode a SPIR number or a fixed
sheet name list. It auto-detects which tabs are real data pages.
"""
import re
import os
import logging
import datetime

from openpyxl.utils import get_column_letter

from .rules import expand_tag
from .db import DATA_DIR
from .reference_data import find_country_mentions
from .workbook_loader import load_workbook

logger = logging.getLogger(__name__)

NON_DATA_SHEETS = {'validation lists', 'cover'}
ANNEXURE_REVIEW_LOG_PATH = os.path.join(DATA_DIR, 'annexure_lookup_review.log')


def _annexure_marker_key(s) -> str:
    key = re.sub(r'[^A-Z0-9]', '', str(s or '').upper())
    # 'Refer Anneure 1' / 'ANNEXURES-1' must key the same as a sheet titled
    # 'Annexure 1' (misspelling, and plural right before the number).
    key = re.sub(r'^ANNEX?URE', 'ANNEXURE', key)
    return re.sub(r'^ANNEXURES(?=\d|$)', 'ANNEXURE', key)


def _annexure_label(s) -> str:
    """Clean display text for an ANNEXURE reference that couldn't be
    resolved -- 'Refer Anneure 1' -> 'Annexure 1' -- kept as the tag so
    the column (and the file) is still processed instead of dropped."""
    text = _ANNEXURE_LEADIN_RE.sub('', str(s or '').strip())
    text = re.sub(r'^ANNEX?URES?\b[\s\-_:.]*', '', text, flags=re.IGNORECASE)
    return f'Annexure {text}'.strip()


_ANNEXURE_LEADIN_RE = re.compile(
    r'^\s*(PLEASE\s+REFER\s+TO|REFER\s+TO|REFER|SEE|AS\s+PER)\s+', re.IGNORECASE)


def _looks_like_annexure_ref(s) -> bool:
    """True if `s` -- after stripping a common lead-in phrase like 'Refer
    to' / 'Refer' / 'See' / 'As per' -- starts with the word 'ANNEXURE'
    (tolerant of the common 'Annexure' -> 'Anneure' misspelling). This is
    deliberately permissive about everything AFTER that word: a plain
    number ('ANNEXURE-1'), or a qualifier before/after it ('ANNEXURE
    (P1)-1'), or anything else a real SPIR uses -- _resolve_annexure_ref
    does the actual sheet matching, dynamically, against whatever
    annexure sheets THIS workbook actually has, so no one naming scheme
    is hard-coded here."""
    text = _ANNEXURE_LEADIN_RE.sub('', str(s or '').strip())
    return bool(re.match(r'ANNEX?URES?\b', text, re.IGNORECASE))


def _resolve_annexure_ref(s, annexure_sheets):
    """Matches a Tag/Model/Serial cell's ANNEXURE reference text to the
    SPECIFIC annexure sheet it means -- however the reference and the
    sheet's own title are actually punctuated or qualified (e.g. a cell
    reading 'ANNEXURE (P1)-1' pointing at a sheet titled 'ANNEXURE (P1)',
    or at one titled 'ANNEXURE (P1)-1' if the workbook has that many, or
    the plain 'ANNEXURE-1' case). Matching is purely dynamic against
    `annexure_sheets`'s own keys -- no specific qualifier like '(P1)' is
    hard-coded -- tried in order:
      1. exact match on the normalized reference text
      2. the LONGEST known annexure sheet key that's a PREFIX of the
         reference (the reference carries extra trailing detail, e.g. a
         sub-item index, beyond the sheet's own base name)
      3. the SHORTEST known annexure sheet key that the reference is a
         prefix of (the sheet's own title carries extra trailing detail
         beyond what was referenced)
    Returns the matching annexure_sheets entry, or None if nothing in
    this workbook matches -- never guessed."""
    text = _ANNEXURE_LEADIN_RE.sub('', str(s or '').strip())
    ref_norm = _annexure_marker_key(text)
    if not ref_norm or not annexure_sheets:
        return None
    if ref_norm in annexure_sheets:
        return annexure_sheets[ref_norm]
    prefixes = [k for k in annexure_sheets if k and ref_norm.startswith(k)]
    if prefixes:
        best = max(prefixes, key=len)
        sub_key = ref_norm[len(best):]
        annexure = annexure_sheets[best]
        return _filter_annexure_by_group(annexure, sub_key) if sub_key else annexure
    supersets = [k for k in annexure_sheets if k.startswith(ref_norm)]
    if supersets:
        return annexure_sheets[min(supersets, key=len)]
    return None


def _is_valid_tag(s) -> bool:
    """A real SPIR TAG NUMBER always mixes letters and digits (e.g.
    '26-PG-908', 'V-8943-A', '68-SP-061') -- never one fixed pattern,
    since different SPIRs format tags differently, but ALWAYS both. A
    letters-only value (a stray label like 'PUMP', 'VALVE', or leftover
    text like 'ANNEXURE' from a reference that didn't resolve) is never a
    real tag, and must never be extracted as one -- checked wherever a
    candidate tag is about to be finalized, on every sheet kind (Main,
    Continuation, Annexure)."""
    text = str(s or '')
    return bool(re.search(r'[A-Za-z]', text)) and bool(re.search(r'[0-9]', text))


def _is_annexure_sheet(ws) -> bool:
    """A sheet named 'ANNEXURE-1', 'ANNEXURE-2', etc. is a REFERENCE sheet,
    never a Tag Number and never a data sheet in its own right -- some
    OTHER sheet's tag cell may point to it (literally 'ANNEXURE-1', or
    worded as 'Refer Annexure 1' -- see _looks_like_annexure_ref), which is a
    pointer meaning "look up the real tag list on the sheet actually
    titled ANNEXURE-1", not a tag itself (see _parse_sheet)."""
    key = _annexure_marker_key(ws.title)
    return key.startswith('ANNEXURE') or key.startswith('ANNEURE')


_SPIR_LABEL_RE = re.compile(r'^(?:\d+\s*)?SPIR\s*(?:NUMBER|NO\.?)\s*:?$', re.IGNORECASE)


def _find_spir_number_label(ws):
    """Main-sheet SPIR number found by its own label rather than by cell
    position, for template variants whose merged header puts the value in
    e.g. V1 instead of Y1. The label may carry the form's field number
    ('25 SPIR NUMBER:') and be spelled 'SPIR NO'. Only the header rows are
    searched, and the value must sit within 6 columns to the label's right
    (so the SPIR-type checkboxes in column AB are never picked up).
    Returns (cell coordinate, value) or None."""
    for r in range(1, 6):
        for c in range(1, min(ws.max_column or 1, 40) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and _SPIR_LABEL_RE.match(v.strip()):
                for cc in range(c + 1, c + 7):
                    cell = ws.cell(row=r, column=cc)
                    if cell.value not in (None, '') and str(cell.value).strip():
                        return cell.coordinate, cell.value
    return None


# ---------------------------------------------------------------------------
# SPIR form structure. Everything below is located from the form's own
# labels/headers; the fixed positions of the known QatarEnergy SPIR form
# (tags C1:F1, model/serial/units rows 4/6/7, item headers in row 6 from
# column G) are only a fallback for sheets that don't carry those labels.
# ---------------------------------------------------------------------------

def _norm_label(v) -> str:
    """Header/label text for matching: upper-case, punctuation -> space,
    whitespace collapsed ('MFR SER'L NO.' -> 'MFR SER L NO')."""
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Z0-9]+', ' ', str(v or '').upper())).strip()


# Row labels in the tag-label column (column B on the known form).
_ROW_LABELS = {
    'tag':    (1, lambda t: t.startswith('EQUIP') or bool(re.search(r'\bTAG\b', t))),
    'model':  (4, lambda t: t.startswith('MFR TYPE') or 'MODEL' in t),
    'serial': (6, lambda t: t.startswith('MFR SER') or 'SERIAL' in t),
    'units':  (7, lambda t: t.startswith('NO OF UNITS') or t.startswith('NUMBER OF UNITS')),
}

# Item-table columns: field -> (known-form column, header matcher).
_ITEM_COLUMNS = {
    'item_no':          (7,  lambda t: t.startswith('ITEM NUMBER') or t in ('ITEM NO', 'ITEM')),
    'qty_fitted':       (8,  lambda t: 'IDENTICAL PARTS' in t),
    'desc':             (9,  lambda t: t.startswith('DESCRIPTION')),
    'dwg_no':           (10, lambda t: t.startswith('DWG') or t.startswith('DRAWING')),
    'mfr_part_no':      (11, lambda t: re.match(r'(MANUFACTURERS?|MFR S?) PART', t) is not None),
    'supplier_part_no': (12, lambda t: re.match(r'SUPPLIERS? PART', t) is not None),
    'material_spec':    (13, lambda t: t.startswith('MATERIAL SPEC')),
    'supplier_ocm':     (15, lambda t: t.startswith('SUPPLIER OCM')),
    'currency':         (21, lambda t: t.startswith('CURRENCY')),
    'unit_price':       (22, lambda t: t.startswith('UNIT PRICE')),
    'delivery_wks':     (23, lambda t: t.startswith('DELIVERY')),
    'uom':              (25, lambda t: t.startswith('UNIT OF MEASURE') or t == 'UOM'),
    'sap_no':           (26, lambda t: t.startswith('SAP NUMBER') or t.startswith('SAP NO')),
    'classification':   (27, lambda t: t.startswith('CLASSIFICATION')),
}


class SpirLayoutError(ValueError):
    """The sheet looks like a SPIR form but its structure can't be mapped
    with confidence -- the file is rejected rather than mis-read."""


def _col_letter(c) -> str:
    return get_column_letter(c)


def _find_row_labels(ws):
    """The tag-label column (B on the form) and the rows of its 'EQUIPMENT
    TAG No' / 'MFR TYPE OR MODEL' / 'MFR SER'L NO.' / 'No. OF UNITS' labels.
    A column counts only if it has the tag label plus at least two of the
    others. Rows not found fall back to the form's 1/4/6/7.
    Returns (label_col or None, {'tag', 'model', 'serial', 'units'})."""
    label_col, rows = None, {}
    for c in range(1, 4):
        found = {}
        for r in range(1, 13):
            t = _norm_label(ws.cell(row=r, column=c).value)
            for key, (_, match) in _ROW_LABELS.items():
                if t and key not in found and match(t):
                    found[key] = r
        if 'tag' in found and len(found) >= 3:
            label_col, rows = c, found
            break
    for key, (fixed_row, _) in _ROW_LABELS.items():
        rows.setdefault(key, fixed_row)
    return label_col, rows


def _settle_tag_row(ws, rows: dict, tag_cols) -> None:
    """The 'EQUIPMENT TAG No' label is a merged cell spanning several rows,
    and vendors sometimes type the tags on its last line (row 3) rather
    than its first. If the label's own row has no tags, use the first row
    below it -- above the MFR TYPE / MODEL row -- that does."""
    def has_tag(r):
        return any(str(ws.cell(row=r, column=c).value or '').strip() not in ('-', '') for c in tag_cols)
    if has_tag(rows['tag']):
        return
    for r in range(rows['tag'] + 1, rows['model']):
        if has_tag(r):
            rows['tag'] = r
            return


def _find_item_header_row(ws, max_row: int = 15):
    """The item table's header row: the first row (within the top of the
    sheet) holding both an ITEM NUMBER and a DESCRIPTION header."""
    last_col = min(ws.max_column or 1, 60)
    for r in range(1, max_row + 1):
        texts = [_norm_label(ws.cell(row=r, column=c).value) for c in range(1, last_col + 1)]
        if (any(_ITEM_COLUMNS['item_no'][1](t) for t in texts if t)
                and any(_ITEM_COLUMNS['desc'][1](t) for t in texts if t)):
            return r
    return None


def _sheet_layout(ws):
    """Where things are on this main sheet. Returns a dict:
      tag_cols      -- columns holding one equipment tag each (C..F on the form)
      rows          -- {'tag', 'model', 'serial', 'units'} row numbers
      cols          -- {item field: column or None}
      header_row / first_item_row / source ('headers' or 'fixed')
    Raises SpirLayoutError when headers are present but ambiguous."""
    label_col, rows = _find_row_labels(ws)
    header_row = _find_item_header_row(ws)
    if header_row is None:
        # No item headers: the known form's fixed positions (older sheets
        # identified by their SPIR number in Y1 rely on this).
        cols = {f: fixed for f, (fixed, _) in _ITEM_COLUMNS.items()}
        _settle_tag_row(ws, rows, range(3, 7))
        return {'tag_cols': list(range(3, 7)), 'rows': rows, 'cols': cols,
                'header_row': None, 'first_item_row': 8, 'source': 'fixed'}

    cols, missing = {}, []
    last_col = min(ws.max_column or 1, 60)
    for field, (_, match) in _ITEM_COLUMNS.items():
        hits = [c for c in range(1, last_col + 1)
                if (t := _norm_label(ws.cell(row=header_row, column=c).value)) and match(t)]
        if len(hits) > 1:
            raise SpirLayoutError(
                f"Sheet '{ws.title.strip()}': the item table has more than one "
                f"'{field.replace('_', ' ').upper()}' column (row {header_row}).")
        cols[field] = hits[0] if hits else None
        if not hits:
            missing.append(field)
    if missing:
        logger.warning('Sheet %r: item columns not found in header row %d, left blank: %s',
                       ws.title, header_row, ', '.join(missing))

    # Tag columns sit between the row-label column and the ITEM NUMBER column.
    first_tag_col = (label_col or 2) + 1
    tag_cols = list(range(first_tag_col, cols['item_no']))
    if not tag_cols:
        raise SpirLayoutError(f"Sheet '{ws.title.strip()}': no equipment tag columns "
                              f"before the ITEM NUMBER column.")
    _settle_tag_row(ws, rows, tag_cols)

    # Items start after the header row, skipping the column-number row
    # ('7, 8, 9, 10A ...') that the form puts right under the headers.
    first = header_row + 1
    for _ in range(2):
        item_v = ws.cell(row=first, column=cols['item_no']).value
        desc_v = ws.cell(row=first, column=cols['desc']).value
        if item_v in (None, '') or isinstance(desc_v, (int, float)):
            first += 1
    return {'tag_cols': tag_cols, 'rows': rows, 'cols': cols,
            'header_row': header_row, 'first_item_row': first, 'source': 'headers'}


def _has_tag(ws, layout) -> bool:
    return any(str(ws.cell(row=layout['rows']['tag'], column=c).value or '').strip() not in ('-', '')
               for c in layout['tag_cols'])


def _data_sheet_rejection(ws):
    """Why `ws` is NOT a SPIR main/data sheet, or None if it is one. A real
    data page has at least one equipment tag in its tag row and a SPIR
    number: next to its 'SPIR NUMBER' label, or in Y1 on the known form.
    A sheet found by label must also carry the main-sheet item table
    (ITEM NUMBER + DESCRIPTION headers) -- that's what tells it apart from
    a Continuation sheet, which also carries a SPIR NUMBER label. An
    ANNEXURE-N-named sheet is always a reference sheet, never a data sheet
    on its own."""
    if ws.title.strip().lower() in NON_DATA_SHEETS:
        return 'supporting sheet (by name)'
    if _is_annexure_sheet(ws):
        return 'annexure reference sheet'
    try:
        layout = _sheet_layout(ws)
    except SpirLayoutError as e:
        if ws['Y1'].value or _find_spir_number_label(ws):
            raise   # clearly a SPIR main sheet: tell the user what's wrong
        return str(e)
    if not _has_tag(ws, layout):
        return 'no equipment tag in the tag row'
    if layout['source'] != 'headers' and _find_inline_label(ws, 'SPIR NUMBER', rows=range(1, 6)):
        # Continuation-sheet structure (inline 'SPIR NUMBER:' label, no item
        # table) -- even when its tags run on into Y1.
        return 'continuation sheet (SPIR NUMBER label but no item table)'
    if ws['Y1'].value:
        return None
    if not _find_spir_number_label(ws):
        return 'no SPIR number (no filled SPIR NUMBER label in rows 1-5, Y1 empty)'
    if layout['source'] != 'headers':
        return 'no item table (ITEM NUMBER / DESCRIPTION headers)'
    return None


def _is_data_sheet(ws) -> bool:
    return _data_sheet_rejection(ws) is None


def _header_value(ws, cell: str, label: str):
    """A main sheet's header field ('EQUIPMENT', 'MANUFACTURER', ...): the
    first value right of its label in the header rows (within 6 columns,
    so the SPIR-type checkboxes in AB are never read), else the known
    form's fixed cell."""
    found = _find_inline_label(ws, label, rows=range(1, 6), max_offset=6)
    if found and found[1] not in (None, ''):
        return found[1]
    return ws[cell].value


def _contact_block(ws, items_end_row: int):
    """The vendor contact block: the cell under the 'MANUFACTURERS/SUPPLIERS
    FOCAL POINT' label below the item table (K32 -> K33 on the known form),
    else the known form's K33."""
    for r in range(items_end_row, min(ws.max_row or 1, items_end_row + 15) + 1):
        for c in range(1, min(ws.max_column or 1, 40) + 1):
            if _norm_label(ws.cell(row=r, column=c).value).startswith('MANUFACTURERS SUPPLIERS FOCAL'):
                return ws.cell(row=r + 1, column=c).value
    return ws['K33'].value


_TAG_KEYWORD_RE = re.compile(r'\bTAG\b')


def _has_tag_keyword(label: str) -> bool:
    """True if the word 'TAG' appears in `label` as its own word -- 'TAG',
    'TAG NO', 'TAG.NO', 'TAG NUMBER', 'VALVE TAG', 'EQUIPMENT TAG NO.',
    'TAGNO' (glued straight onto TAG, no separator) all match. Uses a word
    boundary rather than a plain substring search so a column merely
    CONTAINING the letters t-a-g, like 'VOLTAGE' or 'MONTAGE', is never
    mistaken for a tag column. This is the one keyword search used
    whenever an ANNEXURE-style sheet's tag column has to be found by its
    header text rather than a fixed position -- never one exact header
    string, since real files spell it differently sheet to sheet."""
    return bool(_TAG_KEYWORD_RE.search(label)) or label.startswith('TAG')


def _normalize_tag_key(v) -> str:
    """Makes TAG NUMBER matching robust to stray whitespace and to Excel
    reading the same tag as text in one sheet and as a number in another
    (e.g. '068' vs 68 vs 68.0)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'\s+', ' ', str(v).strip()).upper()


def _build_annexure_index(ws):
    """Reads an ANNEXURE-style sheet -- header row/columns auto-detected so
    layout drift doesn't break the lookup, and the TAG column isn't
    assumed to be literally named 'TAG NO'; a real file was seen labeling
    it 'VALVE TAG' instead. Returns (index, rows):
      - index: {normalized TAG NO: {'serial': ..., 'model': ...}} for
        exact-tag lookups -- used when a Main/Continuation sheet's own
        Model or Serial No cell for an already-real tag is ITSELF an
        ANNEXURE reference rather than a value (e.g. that cell literally
        reads 'Refer Annexure 1') -- see _resolve_annexure_field.
      - rows: ordered [(TAG NO, serial, model, group_key), ...] -- used
        when a SPIR tag column is itself just a pointer to this sheet
        (see _parse_sheet), each row becoming its own physical tag.
        serial/model are None when the sheet has no such column at all
        (some annexure tables list e.g. TAG NOS + TAG DESCRIPTION +
        EQUIPMENT DESCRIPTION instead -- the tag list itself is still the
        point of the sheet, so a missing column must not make the whole
        sheet unusable). group_key is this row's own value (normalized)
        in a column whose header itself starts with this sheet's own
        name (e.g. a sheet 'ANNEXURE-P1' with a column literally headed
        'ANNEXURE-P1 NUMBER') -- present when many tag columns across the
        workbook share ONE big annexure sheet, each pointing at only ITS
        OWN numbered group of rows within it (e.g. 'ANNEXURE (P1)-1'
        meaning just the rows numbered 1, not the whole sheet) -- see
        _filter_annexure_by_group. None when the sheet has no such
        column at all (the simple case: the whole sheet IS one group).
    """
    sheet_key = _annexure_marker_key(ws.title)
    header_row = tag_col = serial_col = model_col = group_col = None
    max_col = min(ws.max_column or 1, 60)
    for r in range(1, 11):
        labels = {}
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v:
                labels[str(v).strip().upper()] = c
        tc = next((c for label, c in labels.items() if _has_tag_keyword(label)), None)
        if tc:
            header_row, tag_col = r, tc
            serial_col = next((c for label, c in labels.items() if label.startswith('SERIAL')), None)
            model_col = next((c for label, c in labels.items()
                               if label.startswith('MODEL') or label.startswith('TYPE')
                               or label.startswith('MFR')), None)
            if sheet_key:
                group_col = next((c for label, c in labels.items()
                                   if _annexure_marker_key(label).startswith(sheet_key)
                                   and _annexure_marker_key(label) != sheet_key), None)
            break
    if not header_row:
        return {}, []

    index = {}
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        tag_val = ws.cell(row=r, column=tag_col).value
        if tag_val in (None, '') or not _is_valid_tag(tag_val):
            continue
        key = _normalize_tag_key(tag_val)
        if not key or key in index:
            continue
        serial_val = ws.cell(row=r, column=serial_col).value if serial_col else None
        if isinstance(serial_val, str):
            serial_val = serial_val.strip()
        model_val = ws.cell(row=r, column=model_col).value if model_col else None
        if isinstance(model_val, str):
            model_val = model_val.strip()
        group_val = ws.cell(row=r, column=group_col).value if group_col else None
        group_key = _normalize_tag_key(group_val) if group_val not in (None, '') else None
        tag_str = tag_val.strip() if isinstance(tag_val, str) else str(tag_val)
        index[key] = {'serial': serial_val, 'model': model_val}
        rows.append((tag_str, serial_val, model_val, group_key))
    return index, rows


def _filter_annexure_by_group(annexure, sub_key: str):
    """When a reference like 'ANNEXURE (P1)-1' points at a big SHARED
    sheet ('ANNEXURE-P1') by number, rather than at a whole sheet of its
    own, and that sheet has its own per-row group/number column (see
    _build_annexure_index), narrows the sheet's rows down to just the
    ones in that specific group -- e.g. only the handful of VALVE TAG
    rows whose own 'ANNEXURE-P1 NUMBER' column reads '1', not all ~1500
    rows on the sheet. Returns the sheet's rows unfiltered when it has no
    group column at all (a plain single-number-per-sheet annexure, e.g.
    'ANNEXURE-1', where the trailing number belongs to the SHEET itself
    and is already handled by the exact-match case in
    _resolve_annexure_ref -- this function is only reached for the
    'extra trailing digits beyond the matched sheet's own name' case).
    Returns None -- a genuine miss, never a silent fall-back to the
    entire sheet -- when the sheet DOES have groups but this specific
    one doesn't exist."""
    rows = annexure['rows']
    if not any(r[3] is not None for r in rows):
        return annexure
    filtered = [r for r in rows if r[3] == sub_key]
    return {'index': annexure['index'], 'rows': filtered} if filtered else None


def _resolve_annexure_field(value, tag, field: str, annexure_sheets):
    """If `value` (a Model or Serial No cell belonging to an already-real
    tag) is ITSELF an ANNEXURE reference -- 'ANNEXURE-1', 'Refer Annexure
    1', 'Refer Anneure 1', 'ANNEXURE (P1)-1', etc. (see
    _looks_like_annexure_ref) -- rather than an actual value, looks up
    THIS specific tag's own `field` ('model' or 'serial') on the
    referenced annexure sheet and returns that instead. Returns `value`
    unchanged when it isn't a reference at all (the normal case). Returns
    None when it IS a reference but the sheet or this tag's row on it
    can't actually be found -- never falls back to the literal reference
    text as if it were the real value."""
    if not _looks_like_annexure_ref(value):
        return value
    annexure = _resolve_annexure_ref(value, annexure_sheets)
    if not annexure:
        return None
    rec = annexure['index'].get(_normalize_tag_key(tag))
    return rec.get(field) if rec else None


def _add_unresolved_annexure_col(tag_cols, col, raw_tag_str, model, sern, qty_units):
    """Tag column for an ANNEXURE reference that yields no real tags (sheet
    missing, or its tag column only reads e.g. 'NA'): kept as one tag
    named e.g. 'Annexure 1' rather than dropped, with the column's own
    model/serial (unless those are references too). When several columns
    point at the same annexure, they're one tag, so their No. of Units
    (row 7) add up on its first column."""
    tag = _annexure_label(raw_tag_str)
    first = next((tc for tc in tag_cols if tc['tag'] == tag), None)
    if (first is not None and isinstance(first['qty_units'], (int, float))
            and isinstance(qty_units, (int, float))):
        first['qty_units'] += qty_units
    tag_cols.append({'col': col, 'tag': tag,
                     'model': None if _looks_like_annexure_ref(model) else model,
                     'sern': None if _looks_like_annexure_ref(sern) else sern,
                     'qty_units': qty_units})


def _log_annexure_miss(spir_no: str, tag: str):
    os.makedirs(DATA_DIR, exist_ok=True)
    try:
        with open(ANNEXURE_REVIEW_LOG_PATH, 'a', encoding='utf-8') as f:
            f.write(f'{datetime.datetime.utcnow().isoformat()}\t{spir_no}\t{tag}\t'
                    f'TAG NUMBER not found in ANNEXURE-1\n')
    except Exception:
        pass


def _find_inline_label(ws, label: str, rows, max_offset: int = 15):
    """Scans the given rows (an int or a range) across every column for a
    cell whose text -- trimmed, trailing ':' dropped, upper-cased -- equals
    `label`, e.g. 'SPIR NUMBER'. Returns (label_col, value): value is the
    first non-blank cell within max_offset columns to the label's right (a
    merged/wide field means the actual value isn't always the very next
    cell), or None for value if the label has no populated cell after it.
    Returns None (not a tuple) if the label isn't present on this sheet at
    all -- position is never hard-coded, so this survives layout drift
    between different SPIR templates/exports."""
    row_range = rows if hasattr(rows, '__iter__') else [rows]
    for r in row_range:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v and str(v).strip().rstrip(':').upper() == label:
                for cc in range(c + 1, c + 1 + max_offset):
                    vv = ws.cell(row=r, column=cc).value
                    if vv not in (None, ''):
                        return c, vv
                return c, None
    return None


def _is_continuation_sheet(ws) -> bool:
    """A CONTINUATION-style sheet (e.g. 'CONTINUATION SHEET', 'CONTINUATION
    SHEET (2)') has no SPIR number in Y1 -- so _is_data_sheet rejects it --
    but it always carries its own inline 'SPIR NUMBER:' label somewhere in
    its top rows. That's the reliable signal that it's a genuine SPIR
    continuation page (extending an adjacent Main Sheet's own tags/items
    to more tags than fit in one page -- see _parse_continuation_sheet)
    rather than an unrelated tab like 'Validation Lists'. Never matched by
    sheet title text -- purely by this structural marker."""
    if ws.title.strip().lower() in NON_DATA_SHEETS or _is_annexure_sheet(ws) or _is_data_sheet(ws):
        return False
    return _find_inline_label(ws, 'SPIR NUMBER', rows=range(1, 6)) is not None


def _item_no_key(v) -> str:
    """Normalizes an ITEM NUMBER for cross-sheet matching (Main Sheet's own
    item numbering vs. a Continuation Sheet's reference to it), robust to
    Excel reading the same number as int in one sheet and float in
    another (e.g. 4 vs 4.0)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse_continuation_sheet(ws, annexure_sheets=None):
    """A CONTINUATION-style sheet defines no items of its own -- it only
    extends items ALREADY defined on its paired Main Sheet (matched by
    ITEM NUMBER, read here from its own S.NO column C) to additional
    physical tags, laid out with the same row convention as a Main Sheet
    (row1=tag, row4=model, row6=serial no, row7=no. of units) but across a
    dynamically wide tag-column range -- however many columns the sheet
    actually uses -- capped wherever its own inline 'SPIR NUMBER:' label
    begins, instead of a fixed C:F. Exactly like a Main Sheet, one of its
    tag cells may itself be an 'ANNEXURE-N' pointer rather than a real tag
    (see _parse_sheet) -- resolved here the same way.
    Returns (tag_cols, item_flags):
      - tag_cols: same shape as _parse_sheet's tag_cols (used by
        parse_spir to extend the paired Main Sheet's own tag_cols, so
        these tags flow into the global tag_order/tag_info exactly like
        any other tag)
      - item_flags: [(item_no, [tag, ...], {tag: qty}), ...] -- one entry
        per item row that has at least one tag flagged on this sheet, to
        be merged into the paired Main Sheet's matching item (by item_no)
        so it picks up the extra tags with the Main Sheet's own
        description/dwg/price/etc. -- never duplicated as a separate,
        description-less item.
    """
    label = _find_inline_label(ws, 'SPIR NUMBER', rows=range(1, 6))
    stop_col = label[0] if label else ws.max_column + 1
    spir_no = label[1] if label else ''
    # Same row labels as a Main Sheet (column B); the column right after
    # them (C on the form) holds this sheet's item numbers from row 8.
    label_col, rows = _find_row_labels(ws)
    item_col = (label_col or 2) + 1
    _settle_tag_row(ws, rows, range(item_col, stop_col))

    tag_cols = []
    for col in range(item_col, stop_col):
        raw_tag = ws.cell(row=rows['tag'], column=col).value
        if raw_tag and str(raw_tag).strip() not in ('-', ''):
            raw_tag_str = str(raw_tag).strip()
            model = ws.cell(row=rows['model'], column=col).value
            sern = ws.cell(row=rows['serial'], column=col).value
            qty_units = ws.cell(row=rows['units'], column=col).value

            if _looks_like_annexure_ref(raw_tag_str):
                # Recognized as an ANNEXURE reference (however it's
                # phrased/spelled/qualified) -- this column is NEVER
                # treated as a literal tag from here on, even if the
                # referenced sheet can't actually be found (logged
                # instead, below).
                annexure = _resolve_annexure_ref(raw_tag_str, annexure_sheets)
                if annexure and annexure['rows']:
                    for tag_str, serial_val, model_val, _group in annexure['rows']:
                        if not _is_valid_tag(tag_str):
                            continue
                        tag_cols.append({'col': col, 'tag': tag_str, 'model': model_val or model,
                                          'sern': serial_val, 'qty_units': 1})
                else:
                    _log_annexure_miss(spir_no, raw_tag_str)
                    _add_unresolved_annexure_col(tag_cols, col, raw_tag_str, model, sern, qty_units)
                continue

            expanded_tags = expand_tag(raw_tag_str)
            n = len(expanded_tags)
            model_parts = _split_paired_values(model, n) or [model] * n
            sern_parts = _split_paired_values(sern, n, try_slash=True) or [sern] * n
            for idx, tag in enumerate(expanded_tags):
                if not _is_valid_tag(tag):
                    continue
                # The Model or Serial No cell for an already-real tag can
                # itself be an ANNEXURE reference ('Refer Annexure 1')
                # instead of a value -- resolve it from that sheet's own
                # per-tag record rather than storing the reference text.
                tag_model = _resolve_annexure_field(model_parts[idx], tag, 'model', annexure_sheets)
                tag_sern = _resolve_annexure_field(sern_parts[idx], tag, 'serial', annexure_sheets)
                tag_cols.append({'col': col, 'tag': tag, 'model': tag_model,
                                  'sern': tag_sern, 'qty_units': qty_units})

    item_flags = []
    r = 8
    while r <= ws.max_row:
        item_no = ws.cell(row=r, column=item_col).value   # C: this sheet's own reference to the Main Sheet's ITEM NUMBER
        if item_no in (None, ''):
            r += 1
            continue
        if not isinstance(item_no, (int, float)):
            # The item-number column is a clean numeric sequence for as
            # long as this is really an item row; a non-numeric value here
            # (seen in real files: a requisition/PO reference in the same
            # column, further down the page) marks the end of the item
            # list, not another item.
            break
        flags = []
        tag_qty = {}
        for tc in tag_cols:
            v = ws.cell(row=r, column=tc['col']).value
            # Several columns can yield the same tag (e.g. C1:F1 all reading
            # 'ANNEXURES-1'); flag it once per item, keeping the first qty.
            if v not in (None, '-', '') and tc['tag'] not in flags:
                flags.append(tc['tag'])
                tag_qty[tc['tag']] = v if isinstance(v, (int, float)) else None
        if flags:
            item_flags.append((item_no, flags, tag_qty))
        r += 1

    return tag_cols, item_flags


_CONTACT_RE = re.compile(r'(?:Tel(?:ephone)?(?:\s*No\.?)?|Phone|Landline)[.:]?\s*([^\n]+)', re.IGNORECASE)


def _parse_vendor_contact(supplier_field: str, contact_block: str):
    """Vendor details, checking BOTH the top-right 'SUPPLIER:' field and the
    bottom vendor/contact block. Never infers -- a field with no explicit
    source in either location is left blank."""
    supplier_field = (supplier_field or '').strip()
    contact_block = contact_block or ''
    combined = f'{supplier_field}\n{contact_block}'

    # VENDOR NAME: prefer the top-right corner; fall back to the first
    # non-blank line of the bottom block (typically the vendor's own name).
    name = supplier_field
    if not name:
        for line in contact_block.splitlines():
            line = line.strip()
            if line:
                name = line
                break

    emails = re.findall(r'[\w\.-]+@[\w\.-]+', combined)
    phones = _CONTACT_RE.findall(combined)
    countries = find_country_mentions(combined)

    return {
        'name': name,
        'email1': emails[0] if emails else '',
        'email2': emails[1] if len(emails) > 1 else '',
        'contact_no': phones[0].strip() if phones else '',
        'country': countries[0] if len(countries) == 1 else '',
    }


SPIR_TYPE_LABELS = {
    'AB2': 'COMMISSIONING SPARES',
    'AB3': 'INITIAL SPARES',
    'AB4': 'NORMAL OPERATING SPARES',
    'AB5': 'LIFE CYCLE SPARES',
}

_REV_PREFIX_RE = re.compile(r'^\s*REV(?:ISION)?\.?\s*', re.IGNORECASE)


def _extract_issue_rev(ws):
    """Finds the 'ISSUE LETTER:' label anywhere on the sheet (position isn't
    hard-coded, so this survives layout drift) and reads the first non-empty
    cell to its right on the same row, e.g. 'Rev A' / 'REV 2' / 'Rev. A' ->
    'A' / '2' / 'A'. Returns None if the label or its value can't be found
    -- callers must leave REV blank and log it rather than guess."""
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v and str(v).strip().rstrip(':').upper() == 'ISSUE LETTER':
                for c in range(cell.column + 1, cell.column + 6):
                    vv = ws.cell(row=cell.row, column=c).value
                    if vv not in (None, ''):
                        rev = _REV_PREFIX_RE.sub('', str(vv).strip()).strip()
                        return rev or None
                return None
    return None


def _split_paired_values(raw, count, try_slash=False):
    """When a tag column's own header cell expands into several physical
    tags (e.g. 'PM-3427A/B'), its Model/Serial cell sometimes lists one
    value per tag too, rather than one shared value -- observed as
    newline-separated ('3G1F2525073135\\n3G1F2525073136') and, for Serial
    Number specifically, also as '/'-separated on one line
    ('4020363.1-1 / 4020363.1-2'), matching the tag cell's own delimiter.
    Returns the values in order if a delimiter's split count matches
    exactly; otherwise None, so the caller falls back to treating `raw` as
    one value shared by every expanded tag -- never a guessed split.

    try_slash defaults to False because a Model Number routinely contains a
    '/' as part of one single designation (e.g. a mounting spec like
    'IMB3/IM1001'), not as a multi-value delimiter -- splitting on it there
    would corrupt a real model into two fake ones. Serial Numbers don't have
    that ambiguity in the samples seen, so callers pass try_slash=True only
    for those."""
    if raw is None or count <= 1:
        return None
    text = str(raw)
    candidates = [[p.strip() for p in text.splitlines() if p.strip() != '']]
    if try_slash:
        candidates.append([p.strip() for p in text.split('/') if p.strip() != ''])
    for parts in candidates:
        if len(parts) == count:
            return parts
    return None


def _parse_sheet(ws, annexure_sheets=None):
    layout = _sheet_layout(ws)
    rows, cols = layout['rows'], layout['cols']
    found = _find_spir_number_label(ws)
    spir_raw = found[1] if found else ws['Y1'].value
    logger.info('SPIR sheet %r: layout from %s (item header row %s, tag columns %s); '
                'SPIR number at %s: %r', ws.title, layout['source'], layout['header_row'],
                [_col_letter(c) for c in layout['tag_cols']], found[0] if found else 'Y1', spir_raw)
    spir_no = str(spir_raw or '').strip().lstrip('_').strip()
    equipment_desc = str(_header_value(ws, 'X2', 'EQUIPMENT') or '').strip()
    manufacturer = str(_header_value(ws, 'Y3', 'MANUFACTURER') or '').strip()
    supplier = str(_header_value(ws, 'W4', 'SUPPLIER') or '').strip()
    # Only trust the SPIR type when exactly one checkbox is selected --
    # 0 or 2+ selected is ambiguous and must not be guessed. The checkboxes
    # are form controls whose captions aren't cell text, so their linked
    # cells (AB2:AB5) are read by position; only a real True counts.
    checked_types = [lbl for k, lbl in SPIR_TYPE_LABELS.items() if ws[k].value is True]
    spir_type = checked_types[0] if len(checked_types) == 1 else ''
    spir_rev = _extract_issue_rev(ws)

    tag_cols = []
    for col in layout['tag_cols']:
        raw_tag = ws.cell(row=rows['tag'], column=col).value
        if raw_tag and str(raw_tag).strip() not in ('-', ''):
            raw_tag_str = str(raw_tag).strip()
            model = ws.cell(row=rows['model'], column=col).value
            sern_from_spir = ws.cell(row=rows['serial'], column=col).value
            qty_units = ws.cell(row=rows['units'], column=col).value

            # A tag cell reading e.g. 'ANNEXURE-1', or phrased as 'Refer
            # Annexure 1' / 'Refer to Annexure-1' / 'ANNEXURE (P1)-1'
            # (also tolerant of the 'Annexure'->'Anneure' misspelling), is
            # a pointer, not a tag: look up the SPECIFIC matching annexure
            # sheet (if this workbook has one) rather than a single fixed
            # one, so a workbook with several (possibly qualified)
            # annexures resolves each tag column to its own correct sheet.
            if _looks_like_annexure_ref(raw_tag_str):
                # Recognized as an ANNEXURE reference -- this column is
                # NEVER treated as a literal tag from here on, even if the
                # referenced sheet can't actually be found (logged instead,
                # below). Each of the referenced sheet's rows is its own
                # physical tag here, each one its own unit (hence
                # qty_units=1 per expanded tag).
                annexure = _resolve_annexure_ref(raw_tag_str, annexure_sheets)
                if annexure and annexure['rows']:
                    for tag_str, serial_val, model_val, _group in annexure['rows']:
                        if not _is_valid_tag(tag_str):
                            continue
                        tag_cols.append({'col': col, 'tag': tag_str, 'model': model_val or model,
                                          'sern': serial_val, 'qty_units': 1})
                else:
                    _log_annexure_miss(spir_no, raw_tag_str)
                    _add_unresolved_annexure_col(tag_cols, col, raw_tag_str, model,
                                                 sern_from_spir, qty_units)
                continue

            expanded_tags = expand_tag(raw_tag_str)
            n = len(expanded_tags)
            model_parts = _split_paired_values(model, n) or [model] * n
            sern_parts = _split_paired_values(sern_from_spir, n, try_slash=True) or [sern_from_spir] * n

            for idx, tag in enumerate(expanded_tags):
                if not _is_valid_tag(tag):
                    continue
                # The Model or Serial No cell for an already-real tag can
                # itself be an ANNEXURE reference ('Refer Annexure 1')
                # instead of a value -- resolve it from that sheet's own
                # per-tag record rather than storing the reference text.
                tag_model = _resolve_annexure_field(model_parts[idx], tag, 'model', annexure_sheets)
                sern_from_row = _resolve_annexure_field(sern_parts[idx], tag, 'serial', annexure_sheets)
                if sern_from_row not in (None, '', '-'):
                    # The SPIR itself states this tag's own Serial Number
                    # directly (row 6, or its per-tag slice of a '/' or
                    # newline-separated cell) -- that's authoritative and
                    # must never be discarded in favor of a lookup table.
                    sern = sern_from_row
                else:
                    sern = None
                tag_cols.append({'col': col, 'tag': tag, 'model': tag_model,
                                  'sern': sern, 'qty_units': qty_units})

    def cell(r, field):
        c = cols[field]
        return ws.cell(row=r, column=c).value if c else None

    items = []
    r = layout['first_item_row']
    while True:
        item_no = cell(r, 'item_no')
        desc = cell(r, 'desc')
        if item_no is None or desc in (None, ''):
            break
        flags = []
        tag_qty = {}   # tag -> actual per-tag BOM quantity, from the tag's own
                        # flag cell (C:F) for this item row -- e.g. a cell
                        # holding 1 or 2, not just a yes/no mark. This is the
                        # real per-BOM-line quantity; it is NOT the same
                        # concept as 'qty_fitted' (column H, "TOTAL NO. OF
                        # IDENTICAL PARTS FITTED", an equipment-level count).
        for tc in tag_cols:
            v = ws.cell(row=r, column=tc['col']).value
            # Several columns can yield the same tag (e.g. C1:F1 all reading
            # 'ANNEXURES-1'); flag it once per item, keeping the first qty.
            if v not in (None, '-', '') and tc['tag'] not in flags:
                flags.append(tc['tag'])
                tag_qty[tc['tag']] = v if isinstance(v, (int, float)) else None
        items.append({
            'item_no': item_no,
            'desc': desc,
            'qty_fitted': cell(r, 'qty_fitted'),   # TOTAL NO. OF IDENTICAL PARTS FITTED
            'dwg_no': cell(r, 'dwg_no'),
            'mfr_part_no': cell(r, 'mfr_part_no'),
            'supplier_part_no': cell(r, 'supplier_part_no'),
            'material_spec': cell(r, 'material_spec'),
            'supplier_ocm': cell(r, 'supplier_ocm'),
            'currency': cell(r, 'currency'),
            'unit_price': cell(r, 'unit_price'),
            'delivery_wks': cell(r, 'delivery_wks'),
            'uom': cell(r, 'uom'),
            'sap_no': cell(r, 'sap_no'),
            'classification': cell(r, 'classification'),
            'flags': flags,
            'tag_qty': tag_qty,
        })
        r += 1

    vendor = _parse_vendor_contact(supplier, _contact_block(ws, r))

    return {
        'sheet_name': ws.title,
        'spir_no': spir_no,
        'spir_rev': spir_rev,
        'equipment_desc': equipment_desc,
        'manufacturer': manufacturer,
        'supplier': supplier,
        'vendor': vendor,
        'spir_type': spir_type,
        'tag_cols': tag_cols,
        'items': items,
    }


def parse_spir(path: str) -> dict:
    """Returns a dict with:
       - sheet_names: ordered list of data-sheet names
       - sheets: {sheet_name: parsed sheet dict}
       - tag_order: unique tags in first-seen order across all sheets
       - tag_info: {tag: {model, sern, qty_units}}  (merged across sheets)
       - spir_no / manufacturer / equipment_desc / spir_type / vendor: from first sheet
    """
    wb = load_workbook(path)   # any supported Excel format, see engine.workbook_loader
    sheet_names = []
    for ws in wb.worksheets:
        reason = _data_sheet_rejection(ws)
        if reason is None:
            sheet_names.append(ws.title)
            logger.info('Detected SPIR sheet: %r', ws.title)
        else:
            logger.info('Not a SPIR main sheet: %r (%s)', ws.title, reason)
    if not sheet_names:
        logger.warning('No SPIR main sheet in %s; checked: %s', os.path.basename(path),
                       '; '.join(f'{ws.title!r}: {_data_sheet_rejection(ws)}' for ws in wb.worksheets))
        raise ValueError('No SPIR main sheet found. Expected a sheet with a SPIR NUMBER, '
                          'equipment tags beside the EQUIPMENT TAG No label, and the '
                          'ITEM NUMBER / DESCRIPTION item table.')

    # Every ANNEXURE-N-named sheet is its own reference sheet, keyed by its
    # own normalized title so a tag cell reading 'ANNEXURE-1' resolves to
    # THAT sheet and one reading 'ANNEXURE-2' resolves to a different one,
    # if the workbook has several.
    annexure_sheets = {}
    for ws in wb.worksheets:
        if _is_annexure_sheet(ws):
            idx, rows = _build_annexure_index(ws)
            annexure_sheets[_annexure_marker_key(ws.title)] = {'index': idx, 'rows': rows}

    sheets = {sn: _parse_sheet(wb[sn], annexure_sheets) for sn in sheet_names}

    # CONTINUATION-style sheets (no SPIR number of their own in Y1, but
    # extending a Main Sheet's tags/items across more columns than fit on
    # one page) are paired to the nearest preceding Main-type data sheet in
    # actual workbook order -- e.g. 'MAIN SHEET' -> 'CONTINUATION SHEET',
    # then 'MAIN SHEET (2)' -> 'CONTINUATION SHEET (2)', and so on for
    # however many Main/Continuation groups the workbook has. Every tag
    # and every item-tag flag a Continuation Sheet contributes is merged
    # straight into its Main Sheet's own tag_cols/items, so it flows
    # through the exact same tag_order/extraction/output logic as any
    # other tag -- nothing downstream needs to know Continuation sheets
    # exist at all.
    current_main = None
    for ws in wb.worksheets:
        if ws.title in sheets:
            current_main = ws.title
        elif current_main is not None and _is_continuation_sheet(ws):
            main_sheet = sheets[current_main]
            cont_tag_cols, item_flags = _parse_continuation_sheet(ws, annexure_sheets)
            main_sheet['tag_cols'].extend(cont_tag_cols)

            item_by_no = {_item_no_key(it['item_no']): it for it in main_sheet['items']}
            for item_no, tags, tag_qty in item_flags:
                it = item_by_no.get(_item_no_key(item_no))
                if it is None:
                    _log_annexure_miss(main_sheet['spir_no'],
                                        f"{ws.title}: item {item_no} has no matching "
                                        f"Main Sheet item to attach tags {tags} to")
                    continue
                for tag in tags:
                    if tag not in it['flags']:
                        it['flags'].append(tag)
                    it['tag_qty'][tag] = tag_qty.get(tag)

    tag_order = []
    tag_info = {}
    for sn in sheet_names:
        for tc in sheets[sn]['tag_cols']:
            tag = tc['tag']
            if tag not in tag_info:
                tag_order.append(tag)
                tag_info[tag] = {'model': tc['model'], 'sern': tc['sern'], 'qty_units': tc['qty_units'], 'sheet': sn}
            elif tag_info[tag]['model'] in (None, '-') and tc['model'] not in (None, '-'):
                tag_info[tag] = {'model': tc['model'], 'sern': tc['sern'], 'qty_units': tc['qty_units'], 'sheet': sn}

    if not tag_order:
        raise ValueError('No valid TAG NUMBER could be extracted from this SPIR file (every tag '
                          'column was blank, unresolved, or letters-only -- see '
                          f'{ANNEXURE_REVIEW_LOG_PATH} for details on any ANNEXURE references '
                          'that could not be resolved).')

    first = sheets[sheet_names[0]]
    return {
        'sheet_names': sheet_names,
        'sheets': sheets,
        'tag_order': tag_order,
        'tag_info': tag_info,
        'spir_no': first['spir_no'],
        'spir_rev': first['spir_rev'],
        'manufacturer': first['manufacturer'],
        'equipment_desc': first['equipment_desc'],
        'supplier': first['supplier'],
        'spir_type': first['spir_type'],
        'vendor': first['vendor'],
    }

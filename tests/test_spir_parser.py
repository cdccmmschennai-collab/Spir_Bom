"""Regression tests for SPIR workbook detection/parsing.

Run from the project root:   python -m unittest discover -s tests -v
To also check every sample SPIR:  set SPIR_SAMPLES_DIR=<samples folder> first.

Output building (Extraction + SAP OUTPUT) runs against a throwaway SQLite
database and log folder, with network lookups (FX rates, manufacturer
country) blocked, so the tests never touch data/ and work offline.
"""
import os
import sys
import shutil
import tempfile
import unittest
import urllib.request
import warnings
from unittest import mock

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from engine import db, parser, sap_output, manufacturer_country  # noqa: E402
from engine.parser import parse_spir  # noqa: E402
from engine.extraction import build_extraction  # noqa: E402
from engine.sap_output import build_sap_output  # noqa: E402

FIXTURES = os.path.join(ROOT, 'tests', 'fixtures')
LABEL_LAYOUT = os.path.join(FIXTURES, 'VEN-4460-DGEN-5-43-0305-1.xlsm')   # SPIR no. in V1, not Y1
STANDARD = os.path.join(FIXTURES, 'VEN-4391-MEWTP-5-43-2003-A.xlsx')      # SPIR no. in Y1


def _offline(*args, **kwargs):
    raise OSError('network disabled in tests')


class SandboxedTestCase(unittest.TestCase):
    """Redirects every file the engine writes (DB, review logs, caches) to a
    temp folder and blocks network access for the duration of each test."""

    def setUp(self):
        warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
        self.tmp = tempfile.mkdtemp()
        patches = [
            mock.patch.object(db, 'DB_PATH', os.path.join(self.tmp, 'test.db')),
            mock.patch.object(parser, 'DATA_DIR', self.tmp),
            mock.patch.object(parser, 'ANNEXURE_REVIEW_LOG_PATH', os.path.join(self.tmp, 'annexure.log')),
            mock.patch.object(sap_output, 'DATA_DIR', self.tmp),
            mock.patch.object(sap_output, 'REVIEW_LOG_PATH', os.path.join(self.tmp, 'output.log')),
            mock.patch.object(manufacturer_country, 'DATA_DIR', self.tmp),
            mock.patch.object(manufacturer_country, 'CACHE_PATH', os.path.join(self.tmp, 'mc_cache.json')),
            mock.patch.object(manufacturer_country, 'REVIEW_LOG_PATH', os.path.join(self.tmp, 'mc.log')),
            mock.patch.object(urllib.request, 'urlopen', _offline),
        ]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        db.init_db()


class LabelPositionedSpirTest(SandboxedTestCase):
    """VEN-4460-DGEN-5-43-0305-1.xlsm: SPIR number next to '25 SPIR NUMBER:'
    in V1 (Y1 empty), 3 main sheets + 2 'Cont Sheet' continuation sheets,
    tag cells reading 'ANNEXURES-1' / 'ANNEXURES-2'."""

    @classmethod
    def setUpClass(cls):
        warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

    def test_workbook_is_accepted(self):
        parsed = parse_spir(LABEL_LAYOUT)   # must not raise 'No SPIR ... sheet found'
        self.assertEqual(parsed['sheet_names'], ['Main Sheet-1', 'Main sheet-2', 'Main Sheet-3'])

    def test_spir_number_and_header_fields(self):
        parsed = parse_spir(LABEL_LAYOUT)
        # Raw cell text is kept as-is, exactly like the standard (Y1) layout.
        self.assertTrue(parsed['spir_no'].startswith('VEN-4460-DGEN-5-43-0305'))
        self.assertIn('Rev.1', parsed['spir_no'])
        self.assertEqual(parsed['spir_rev'], '1')
        self.assertEqual(parsed['spir_type'], 'NORMAL OPERATING SPARES')
        self.assertEqual(parsed['equipment_desc'], 'BALL, GATE, GLOBE VALVES')
        self.assertEqual(parsed['manufacturer'], 'JC VALVES')
        self.assertEqual(parsed['supplier'], 'JC VALVES')

    def test_continuation_sheets_are_not_main_sheets(self):
        wb = openpyxl.load_workbook(LABEL_LAYOUT, data_only=True)
        for title in ('Cont Sheet-1 ', 'Cont Sheet -3'):
            self.assertFalse(parser._is_data_sheet(wb[title]), title)
            self.assertTrue(parser._is_continuation_sheet(wb[title]), title)

    def test_annexure_references(self):
        parsed = parse_spir(LABEL_LAYOUT)
        # 'ANNEXURES-1' -> sheet 'Annexure -1 ': its real tags are used.
        self.assertIn('13-BV-5054', parsed['tag_order'])
        # 'ANNEXURES-2' -> sheet 'Annexure -2 ' lists only 'NA' tags: kept as
        # one 'Annexure 2' tag; C7+D7+E7 units (6+5+4) = its 15 rows.
        self.assertIn('Annexure 2', parsed['tag_order'])
        self.assertEqual(parsed['tag_info']['Annexure 2']['qty_units'], 15)
        self.assertEqual(len(parsed['tag_order']), 87)
        for sheet in parsed['sheets'].values():
            for it in sheet['items']:
                self.assertEqual(len(it['flags']), len(set(it['flags'])), it['item_no'])

    def test_outputs_build(self):
        parsed = parse_spir(LABEL_LAYOUT)
        ext = os.path.join(self.tmp, 'ext.xlsx')
        out = os.path.join(self.tmp, 'out.xlsx')
        build_extraction(parsed, ext)
        build_sap_output(parsed, out, spir_filename='VEN-4460-DGEN-5-43-0305-1', job_id='test0305')
        ext_values = {c.value for row in openpyxl.load_workbook(ext).active.iter_rows() for c in row}
        self.assertIn('13-BV-5054', ext_values)
        self.assertIn('Annexure 2', ext_values)
        self.assertGreater(openpyxl.load_workbook(out).active.max_row, 1)


class StandardSpirTest(SandboxedTestCase):
    """Standard template (SPIR number in Y1) must parse exactly as before."""

    def test_standard_layout_unchanged(self):
        parsed = parse_spir(STANDARD)
        self.assertEqual(parsed['sheet_names'], ['MAIN SHEET'])
        self.assertEqual(parsed['spir_no'], 'VEN-4391-MEWTP-5-43-2003')
        self.assertEqual(len(parsed['tag_order']), 18)

    def test_standard_outputs_build(self):
        parsed = parse_spir(STANDARD)
        build_extraction(parsed, os.path.join(self.tmp, 'ext.xlsx'))
        build_sap_output(parsed, os.path.join(self.tmp, 'out.xlsx'),
                         spir_filename='VEN-4391-MEWTP-5-43-2003-A', job_id='test2003')


class DetectionHelpersTest(unittest.TestCase):

    def test_annexure_key_normalization(self):
        key = parser._annexure_marker_key
        self.assertEqual(key('ANNEXURES-1'), key('Annexure -1 '))
        self.assertEqual(key('Refer Anneure 1'.split(' ', 1)[1]), key('Annexure 1'))
        self.assertEqual(key('ANNEXURE (P1)-1'), 'ANNEXUREP11')   # unchanged
        self.assertTrue(parser._looks_like_annexure_ref('ANNEXURES-2'))
        self.assertEqual(parser._annexure_label('ANNEXURES-2'), 'Annexure 2')
        self.assertEqual(parser._annexure_label('Refer Anneure 1'), 'Annexure 1')

    def test_non_spir_workbook_rejected_with_clear_error(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        path = os.path.join(tmp, 'not_spir.xlsx')
        wb = openpyxl.Workbook()
        wb.active['A1'] = 'Just a list'
        wb.active['C1'] = 'PUMP-01'
        wb.active['U1'] = 'SPIR NUMBER:'   # label alone is not enough
        wb.active['V1'] = 'VEN-1'
        wb.save(path)
        with self.assertRaisesRegex(ValueError, 'No SPIR main sheet found'):
            parse_spir(path)


def _snapshot(parsed: dict) -> dict:
    """Comparable form of a parse result (same shape as expected_samples.json)."""
    out = {k: parsed[k] for k in ('sheet_names', 'tag_order', 'spir_no', 'spir_rev', 'manufacturer',
                                  'equipment_desc', 'supplier', 'spir_type', 'vendor')}
    out['tag_info'] = {t: {k: str(v) for k, v in i.items()} for t, i in parsed['tag_info'].items()}
    out['items'] = {sn: [{k: (v if k == 'flags' else str(v)) for k, v in it.items()} for it in s['items']]
                    for sn, s in parsed['sheets'].items()}
    out['tag_cols'] = {sn: [{k: str(v) for k, v in tc.items()} for tc in s['tag_cols']]
                       for sn, s in parsed['sheets'].items()}
    return out


class AllSampleSpirsTest(unittest.TestCase):
    """Every sample SPIR must parse exactly as recorded in
    fixtures/expected_samples.json. The samples are too large for the repo,
    so point SPIR_SAMPLES_DIR at the samples folder (the one holding
    abi_samples/, Arthi/, ...) to run this; otherwise it is skipped."""

    def test_samples_match_expected(self):
        import json
        samples_dir = os.environ.get('SPIR_SAMPLES_DIR')
        if not samples_dir or not os.path.isdir(samples_dir):
            self.skipTest('set SPIR_SAMPLES_DIR to the sample SPIR folder')
        warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
        with open(os.path.join(FIXTURES, 'expected_samples.json'), encoding='utf-8') as f:
            expected = json.load(f)
        checked = 0
        with mock.patch.object(parser, '_log_annexure_miss', lambda *a: None):
            for rel, exp in sorted(expected.items()):
                path = os.path.join(samples_dir, *rel.split('/'))
                if not os.path.exists(path):
                    continue
                with self.subTest(sample=rel):
                    try:
                        got = _snapshot(parse_spir(path))
                    except ValueError as e:
                        got = {'ERROR': str(e)}
                    if 'ERROR' in exp:
                        self.assertIn('ERROR', got)
                    else:
                        self.assertEqual(got, exp)
                checked += 1
        self.assertGreater(checked, 0, 'no sample files found under SPIR_SAMPLES_DIR')


class FormVariationTest(unittest.TestCase):
    """Altered copies of a real SPIR, simulating future revisions of the
    form: the parser must follow the labels, or refuse the file -- never
    read the wrong column."""

    def setUp(self):
        warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        log = mock.patch.object(parser, '_log_annexure_miss', lambda *a: None)
        log.start()
        self.addCleanup(log.stop)
        self.original = parse_spir(STANDARD)

    def _variant(self, change):
        wb = openpyxl.load_workbook(STANDARD)
        change(wb['MAIN SHEET'])
        path = os.path.join(self.tmp, 'variant.xlsx')
        wb.save(path)
        return path

    @staticmethod
    def _item_fields(parsed):
        return [{k: v for k, v in it.items() if k not in ('flags', 'tag_qty')}
                for it in parsed['sheets']['MAIN SHEET']['items']]

    def test_item_columns_reordered(self):
        # A revised form moving columns around: CURRENCY <-> MIN/MAX STOCK
        # (U <-> X) and SAP NUMBER <-> CLASSIFICATION (Z <-> AA), header
        # and data together. Every field must still come from its header.
        last_item_row = 7 + len(self.original['sheets']['MAIN SHEET']['items'])

        def swap(ws):
            for a, b in ((21, 24), (26, 27)):
                for r in range(6, last_item_row + 1):
                    ca, cb = ws.cell(row=r, column=a), ws.cell(row=r, column=b)
                    ca.value, cb.value = cb.value, ca.value
        parsed = parse_spir(self._variant(swap))
        self.assertEqual(parsed['tag_order'], self.original['tag_order'])
        self.assertEqual(self._item_fields(parsed), self._item_fields(self.original))

    def test_tags_on_last_line_of_merged_label(self):
        def move_tags_down(ws):
            for rng in [m for m in ws.merged_cells.ranges if m.min_row <= 3 and m.max_col >= 3 and m.min_col <= 6]:
                ws.unmerge_cells(str(rng))
            for c in range(3, 7):
                ws.cell(row=3, column=c).value = ws.cell(row=1, column=c).value
                ws.cell(row=1, column=c).value = None
        parsed = parse_spir(self._variant(move_tags_down))
        self.assertEqual(parsed['tag_order'], self.original['tag_order'])

    def test_duplicate_header_is_rejected(self):
        def duplicate_currency(ws):
            ws.cell(row=6, column=24).value = 'CURRENCY'   # X6 (normally MIN/MAX STOCK)
        with self.assertRaisesRegex(ValueError, "more than one 'CURRENCY' column"):
            parse_spir(self._variant(duplicate_currency))

    def test_missing_optional_header_left_blank(self):
        def drop_sap_header(ws):
            ws.cell(row=6, column=26).value = 'REMARKS'    # Z6 was SAP NUMBER
        parsed = parse_spir(self._variant(drop_sap_header))
        self.assertTrue(all(it['sap_no'] is None for it in parsed['sheets']['MAIN SHEET']['items']))
        self.assertEqual([it['desc'] for it in parsed['sheets']['MAIN SHEET']['items']],
                         [it['desc'] for it in self.original['sheets']['MAIN SHEET']['items']])


if __name__ == '__main__':
    unittest.main()
